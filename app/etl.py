"""Excel parsing, normalization, rejected-row audit and deterministic deduplication."""
from __future__ import annotations
from dataclasses import dataclass, asdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import posixpath
import re
import unicodedata
import zipfile
from xml.etree import ElementTree
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.database import ImportBatch, LedgerRow

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 5000
MAX_SHEETS = 12
MAX_ISSUES = 250
HEADER_SCAN_ROWS = 12
MAX_COLUMNS = 60
MAX_VISITED_ROWS = (2 * MAX_ROWS) + (MAX_SHEETS * HEADER_SCAN_ROWS)
MAX_VISITED_CELLS = MAX_VISITED_ROWS * MAX_COLUMNS
MAX_XML_NONEMPTY_ROWS = MAX_ROWS + (MAX_SHEETS * HEADER_SCAN_ROWS)

ALIASES = {
    "document_no": {"ma chung tu", "ma hoa don", "so chung tu", "so hoa don", "invoice", "invoice id", "invoice number", "document id", "ma ct"},
    "txn_date": {"ngay", "ngay giao dich", "ngay chung tu", "date", "transaction date", "ngay phat sinh"},
    "department": {"bo phan", "phong ban", "department", "team", "don vi"},
    "counterparty": {"doi tac", "khach hang", "nha cung cap", "partner", "customer", "vendor", "counterparty"},
    "entry_type": {"loai", "loai giao dich", "type", "entry type", "nhom", "loai thu chi"},
    "amount": {"so tien", "gia tri", "amount", "value", "thanh tien", "so tien vnd"},
    "status": {"trang thai", "status", "tinh trang"},
}
REQUIRED = {"document_no", "txn_date", "department", "entry_type", "amount"}


def normalized_header(s: object) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def normalized_identity(value: object) -> str:
    """Normalize case/spacing while retaining Unicode and punctuation identity."""
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().strip()
    return re.sub(r"\s+", " ", text)


def parse_date(v: object) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 30000 <= float(v) <= 80000:
        d = from_excel(v)
        return d.date() if isinstance(d, datetime) else d
    raw = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError("Invalid date: expected dd/mm/yyyy or yyyy-mm-dd")


def parse_amount(v: object) -> Decimal:
    if isinstance(v, bool) or v is None:
        raise ValueError("Missing amount")
    if isinstance(v, (float, int, Decimal)):
        amt = Decimal(str(v))
    else:
        s = str(v).strip().replace("\u00a0", "").replace(" ", "")
        s = re.sub(r"(?i)(vnd|vnđ|đ|₫)", "", s)
        if not re.fullmatch(r"[+-]?[0-9][0-9.,]*", s):
            raise ValueError("Invalid monetary value")
        if "," in s and "." in s:
            decimal_sep = "," if s.rfind(",") > s.rfind(".") else "."
            thousands_sep = "." if decimal_sep == "," else ","
            s = s.replace(thousands_sep, "").replace(decimal_sep, ".")
        elif "," in s or "." in s:
            sep = "," if "," in s else "."
            p = s.split(sep)
            if len(p) > 2 or (len(p) == 2 and len(p[1]) == 3):
                if not all(len(x) == 3 for x in p[1:]):
                    raise ValueError("Ambiguous thousands separators")
                s = "".join(p)
            elif len(p) == 2:
                s = s.replace(sep, ".")
        try:
            amt = Decimal(s)
        except InvalidOperation as exc:
            raise ValueError("Invalid amount") from exc
    if not amt.is_finite() or amt <= 0 or amt > Decimal("99999999999999"):
        raise ValueError("Amount must be positive and within allowed range")
    return amt.quantize(Decimal("0.01"))


def normalize_type(v: object) -> str:
    t = normalized_header(v)
    categories = {
        "REVENUE": {"thu", "doanh thu", "revenue", "income", "ban hang", "sales"},
        "EXPENSE": {"chi", "chi phi", "expense", "cost", "mua hang"},
        "RECEIVABLE": {"phai thu", "cong no phai thu", "receivable", "ar"},
        "PAYABLE": {"phai tra", "cong no phai tra", "payable", "ap"},
    }
    for key, values in categories.items():
        if t in values:
            return key
    raise ValueError(f"Unknown entry type: {str(v)[:35]}")


def normalized_status(v: object) -> str:
    t = normalized_header(v) or "unknown"
    return {"da thanh toan":"paid", "paid":"paid", "hoan tat":"paid", "chua thanh toan":"unpaid", "unpaid":"unpaid", "cho xu ly":"pending", "pending":"pending"}.get(t,"unknown")


def trim(v: object, maximum: int = 100) -> str:
    # Formula cells are read only from their cached values (data_only=True).
    return str(v or "").strip()[:maximum]


def fingerprint_for(
    document_no: object,
    txn_date: date,
    department: object,
    counterparty: object,
    entry_type: object,
    amount: Decimal,
) -> str:
    """Versioned demo identity used for idempotency, including both parties."""
    base = "|".join([
        "v2",
        normalized_identity(document_no),
        txn_date.isoformat(),
        normalized_identity(department),
        normalized_identity(counterparty),
        str(entry_type),
        str(amount),
    ])
    return sha256(base.encode("utf8")).hexdigest()


def legacy_fingerprints_for(
    document_no: object,
    txn_date: date,
    department: object,
    counterparty: object,
    entry_type: object,
    amount: Decimal,
) -> set[str]:
    """Fingerprints written by the two pre-v2 local demo implementations."""
    common = [
        normalized_header(document_no), txn_date.isoformat(),
        normalized_header(department),
    ]
    old_without_party = "|".join([*common, str(entry_type), str(amount)])
    old_with_party = "|".join([
        *common, normalized_header(counterparty), str(entry_type), str(amount),
    ])
    return {
        sha256(old_without_party.encode("utf8")).hexdigest(),
        sha256(old_with_party.encode("utf8")).hexdigest(),
    }


def preflight_worksheet_xml(archive: zipfile.ZipFile, contents: list[zipfile.ZipInfo]) -> None:
    """Bound the worksheet XML OpenPyXL must parse, including ignored columns."""
    items_by_name = {item.filename: item for item in contents}

    def resolve_part(base_part: str, target: str) -> str:
        if target.startswith("/"):
            resolved = posixpath.normpath(target.lstrip("/"))
        else:
            resolved = posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))
        if resolved == ".." or resolved.startswith("../"):
            raise ValueError("Invalid XLSX relationship target")
        return resolved

    def relationships_path(part: str) -> str:
        return posixpath.join(
            posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels",
        )

    try:
        content_types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        override_types = {
            node.attrib.get("PartName", "").lstrip("/"): node.attrib.get("ContentType", "")
            for node in content_types
            if node.tag.rsplit("}", 1)[-1] == "Override"
        }
        default_types = {
            node.attrib.get("Extension", "").lower(): node.attrib.get("ContentType", "")
            for node in content_types
            if node.tag.rsplit("}", 1)[-1] == "Default"
        }

        root_rels = ElementTree.fromstring(archive.read("_rels/.rels"))
        office_rel = next(
            (
                node for node in root_rels
                if node.tag.rsplit("}", 1)[-1] == "Relationship"
                and node.attrib.get("Type", "").endswith("/officeDocument")
            ),
            None,
        )
        if office_rel is None or office_rel.attrib.get("TargetMode") == "External":
            raise ValueError("Invalid XLSX workbook relationship")
        workbook_name = resolve_part("", office_rel.attrib.get("Target", ""))
        workbook = ElementTree.fromstring(archive.read(workbook_name))
        workbook_rels = ElementTree.fromstring(
            archive.read(relationships_path(workbook_name))
        )
        rel_by_id = {
            node.attrib.get("Id", ""): node
            for node in workbook_rels
            if node.tag.rsplit("}", 1)[-1] == "Relationship"
        }
        sheet_ids = [
            next(
                (
                    value for key, value in node.attrib.items()
                    if key.rsplit("}", 1)[-1] == "id"
                ),
                "",
            )
            for node in workbook.iter()
            if node.tag.rsplit("}", 1)[-1] == "sheet"
        ]
        worksheet_names: set[str] = set()
        for rel_id in sheet_ids:
            relation = rel_by_id.get(rel_id)
            if relation is None or relation.attrib.get("TargetMode") == "External":
                raise ValueError("Invalid XLSX worksheet relationship")
            name = resolve_part(workbook_name, relation.attrib.get("Target", ""))
            content_type = override_types.get(
                name,
                default_types.get(name.rsplit(".", 1)[-1].lower(), ""),
            )
            if content_type.endswith(".chartsheet+xml"):
                continue
            if not content_type.endswith(".worksheet+xml"):
                raise ValueError("Invalid XLSX worksheet content type")
            if name not in items_by_name:
                raise ValueError("Invalid XLSX worksheet reference")
            worksheet_names.add(name)
        worksheet_items = [items_by_name[name] for name in worksheet_names]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unreadable XLSX package metadata") from exc
    xml_rows = 0
    xml_cells = 0
    xml_nonempty_rows = 0
    try:
        for item in worksheet_items:
            row_has_value = False
            with archive.open(item) as source:
                for event, element in ElementTree.iterparse(source, events=("start", "end")):
                    name = element.tag.rsplit("}", 1)[-1]
                    if event == "start" and name == "row":
                        row_has_value = False
                        continue
                    if event != "end":
                        continue
                    if name == "row":
                        xml_rows += 1
                        if xml_rows > MAX_VISITED_ROWS:
                            raise ValueError("Workbook exceeds worksheet traversal safety limits")
                        if row_has_value:
                            xml_nonempty_rows += 1
                            if xml_nonempty_rows > MAX_XML_NONEMPTY_ROWS:
                                raise ValueError("Workbook exceeds worksheet traversal safety limits")
                    elif name == "c":
                        xml_cells += 1
                        if xml_cells > MAX_VISITED_CELLS:
                            raise ValueError("Workbook exceeds worksheet traversal safety limits")
                        row_has_value = row_has_value or any(
                            child.tag.rsplit("}", 1)[-1] in {"v", "f", "is"}
                            for child in element
                        )
                    element.clear()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unreadable worksheet XML") from exc


def find_header(rows: list[tuple]) -> tuple[int,dict[str,int]] | None:
    for idx,row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cols = {}
        for col,val in enumerate(row[:MAX_COLUMNS]):
            key = normalized_header(val)
            for target, aliases in ALIASES.items():
                if key in aliases and target not in cols:
                    cols[target] = col
        if REQUIRED <= cols.keys():
            return idx,cols
    return None


def row_value(row: tuple, col: int | None) -> object:
    return row[col] if col is not None and col < len(row) else None


def extract(data: bytes):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("File must be nonempty and no larger than 8MB")
    if not zipfile.is_zipfile(BytesIO(data)):
        raise ValueError("Invalid XLSX archive")
    # Compressed size is insufficient protection against ZIP bombs.
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            contents = archive.infolist()
            if len(contents) > 400 or sum(item.file_size for item in contents) > 64 * 1024 * 1024:
                raise ValueError("Workbook exceeds unpacked safety limits")
            if any(item.file_size > 32 * 1024 * 1024 for item in contents):
                raise ValueError("Workbook contains an oversized member")
            preflight_worksheet_xml(archive, contents)
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid XLSX archive") from exc
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("Unreadable Excel workbook") from exc
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise ValueError("Too many worksheets")
        seen = 0
        visited_rows = 0
        visited_cells = 0
        for ws in wb.worksheets:
            # Read-only worksheets trust attacker-controlled dimension metadata by
            # default. Reset it and enforce our own traversal budgets instead.
            ws.reset_dimensions()
            rows = ws.iter_rows(max_col=MAX_COLUMNS, values_only=True)

            def next_checked():
                nonlocal visited_rows, visited_cells
                row = next(rows, None)
                if row is None:
                    return None
                visited_rows += 1
                visited_cells += len(row)
                if visited_rows > MAX_VISITED_ROWS or visited_cells > MAX_VISITED_CELLS:
                    raise ValueError("Workbook exceeds worksheet traversal safety limits")
                return row

            pre = []
            for _ in range(HEADER_SCAN_ROWS):
                row = next_checked()
                if row is None: break
                pre.append(row)
            header = find_header(pre)
            if not header:
                yield (ws.title, 0, None, "Skipped worksheet without required columns")
                continue
            hidx,columns = header
            from itertools import chain
            remaining = iter(lambda: next_checked(), None)
            for rownum, row in enumerate(chain(pre[hidx+1:],remaining), start=hidx+2):
                if not any(x is not None and str(x).strip() for x in row):
                    continue
                seen+=1
                if seen > MAX_ROWS:
                    raise ValueError(f"Exceeded {MAX_ROWS} nonempty rows")
                try:
                    obj = {
                        "document_no": trim(row_value(row,columns.get("document_no"))),
                        "txn_date": parse_date(row_value(row,columns.get("txn_date"))),
                        "department": trim(row_value(row,columns.get("department"))),
                        "counterparty": trim(row_value(row,columns.get("counterparty")),200) or "Unknown",
                        "entry_type": normalize_type(row_value(row,columns.get("entry_type"))),
                        "amount": parse_amount(row_value(row,columns.get("amount"))),
                        "status": normalized_status(row_value(row,columns.get("status"))),
                    }
                    if not obj["document_no"] or not obj["department"]:
                        raise ValueError("Missing document number or department")
                    obj["fingerprint"] = fingerprint_for(
                        obj["document_no"], obj["txn_date"], obj["department"],
                        obj["counterparty"], obj["entry_type"], obj["amount"],
                    )
                    yield (ws.title,rownum,obj,None)
                except ValueError as exc:
                    yield (ws.title,rownum,None,str(exc))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unreadable worksheet content") from exc
    finally:
        wb.close()


def import_excel(db: Session, filename: str, data: bytes) -> dict:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Only .xlsx is accepted (no macros)")
    # Parse before opening a DB transaction: unreadable files never create batches.
    entries = list(extract(data))
    batch = ImportBatch(filename=filename.split("/")[-1].split("\\")[-1][:255])
    db.add(batch)
    db.flush()
    issues = []
    # Keep old local databases idempotent without rewriting user data. New
    # rows use the Unicode-safe v2 identity; legacy hashes are checked only
    # against hashes that already exist in the database.
    existing = db.execute(select(
        LedgerRow.document_no, LedgerRow.txn_date, LedgerRow.department,
        LedgerRow.counterparty, LedgerRow.entry_type, LedgerRow.amount,
        LedgerRow.fingerprint,
    )).all()
    stored_fingerprints = {row.fingerprint for row in existing}
    fingerprints = set(stored_fingerprints)
    current_identities = {fingerprint_for(*row[:6]) for row in existing}
    stored_parties: dict[str, set[str]] = {}
    for row in existing:
        stored_parties.setdefault(row.fingerprint, set()).add(normalized_identity(row.counterparty))
    for sheet,rownum,obj,error in entries:
        if rownum == 0:
            if len(issues) < MAX_ISSUES:
                issues.append({"sheet":sheet,"row":0,"reason":error})
            continue
        batch.scanned+=1
        if error:
            batch.rejected+=1
            if len(issues)<MAX_ISSUES: issues.append({"sheet":sheet,"row":rownum,"reason":error})
            continue
        legacy = legacy_fingerprints_for(
            obj["document_no"], obj["txn_date"], obj["department"],
            obj["counterparty"], obj["entry_type"], obj["amount"],
        )
        same_legacy_party = any(
            normalized_identity(obj["counterparty"]) in stored_parties.get(old_hash, set())
            for old_hash in legacy
        )
        if (
            obj["fingerprint"] in fingerprints
            or obj["fingerprint"] in current_identities
            or same_legacy_party
        ):
            batch.duplicates+=1
            if len(issues)<MAX_ISSUES: issues.append({"sheet":sheet,"row":rownum,"reason":"Duplicate record (fingerprint)"})
            continue
        fingerprints.add(obj["fingerprint"])
        db.add(LedgerRow(batch_id=batch.id,source_sheet=sheet,source_row=rownum,**obj))
        batch.accepted+=1
    batch.issues_json=json.dumps(issues,ensure_ascii=False)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise ValueError("Concurrent import conflict; retry")
    return {"batch_id":batch.id,"filename":batch.filename,"scanned":batch.scanned,"accepted":batch.accepted,"rejected":batch.rejected,"duplicates":batch.duplicates,"issues":issues,"issues_truncated":len(issues)>=MAX_ISSUES}
