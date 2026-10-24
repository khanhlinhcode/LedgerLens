from io import BytesIO
from datetime import date
from decimal import Decimal
from zipfile import ZipFile, ZIP_DEFLATED
import asyncio
from openpyxl import Workbook
from openpyxl import load_workbook
from app.database import LedgerRow
from app.etl import extract
from app.main import EarlyAPIRequestGate
from app.reports import summarize, workbook_bytes

def test_no_formula_injection_in_report():
    item=LedgerRow(document_no='X-1',txn_date=date(2026,10,1),department='Test',counterparty='=2+2',
       entry_type='REVENUE',amount=Decimal('2000.00'),status='paid',source_sheet='Sheet1',source_row=2,fingerprint='x',batch_id='abc')
    doc=workbook_bytes([item],summarize([item]))
    wb=load_workbook(BytesIO(doc),read_only=True,data_only=False)
    cell=wb['Normalized Data']['E5']
    assert cell.value=='=2+2'
    assert cell.data_type!='f'
    wb.close()


def test_exact_decimal_summary_and_export_lineage():
    item=LedgerRow(document_no='X-2',txn_date=date(2026,10,1),department='Test',counterparty='Customer',
       entry_type='REVENUE',amount=Decimal('99999999999999.99'),status='paid',source_sheet='Sheet1',source_row=7,fingerprint='y',batch_id='batch-1')
    summary=summarize([item])
    assert summary['totals']['REVENUE']=='99999999999999.99'
    doc=workbook_bytes([item],summary)
    wb=load_workbook(BytesIO(doc),read_only=True,data_only=False)
    headers=list(next(wb['Normalized Data'].iter_rows(min_row=4,max_row=4,values_only=True)))
    values=list(next(wb['Normalized Data'].iter_rows(min_row=5,max_row=5,values_only=True)))
    assert 'SOURCE ROW' in headers and 'SOURCE BATCH ID' in headers
    assert values[headers.index('AMOUNT VND')]=='99999999999999.99'
    assert values[headers.index('SOURCE ROW')]==7
    assert values[headers.index('SOURCE BATCH ID')]=='batch-1'
    wb.close()


def test_inflated_worksheet_dimension_hits_traversal_limit():
    original=BytesIO()
    wb=Workbook()
    ws=wb.active
    ws.append(['Invoice','Date','Department','Customer','Type','Amount','Status'])
    ws.append(['A-1','01/10/2026','Sales','Customer','Revenue','1000','paid'])
    wb.save(original)
    source=ZipFile(BytesIO(original.getvalue()))
    crafted=BytesIO()
    with ZipFile(crafted,'w',ZIP_DEFLATED) as target:
        for item in source.infolist():
            data=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                data=data.replace(
                    b'</sheetData>',
                    b'<row r="1048576"><c r="A1048576" t="inlineStr"><is><t>x</t></is></c></row></sheetData>',
                )
            target.writestr(item,data)
    source.close()
    try:
        list(extract(crafted.getvalue()))
        assert False, 'sparse worksheet should be rejected'
    except ValueError as exc:
        assert 'traversal safety limits' in str(exc)


def test_ancillary_data_beyond_supported_columns_does_not_break_ledger_import():
    out=BytesIO()
    wb=Workbook()
    ws=wb.active
    ws.append(['Invoice','Date','Department','Customer','Type','Amount','Status'])
    ws.append(['A-1','01/10/2026','Sales','Customer','Revenue','1000','paid'])
    ws['BI1']='ancillary note'
    wb.save(out)
    events=list(extract(out.getvalue()))
    assert len(events)==1
    assert events[0][2]['document_no']=='A-1'


def test_worksheet_preflight_follows_content_type_not_default_path():
    original=BytesIO()
    wb=Workbook()
    ws=wb.active
    ws.append(['Invoice','Date','Department','Customer','Type','Amount','Status'])
    ws.append(['A-2','01/10/2026','Sales','Customer','Revenue','1000','paid'])
    wb.save(original)
    source=ZipFile(BytesIO(original.getvalue()))
    crafted=BytesIO()
    with ZipFile(crafted,'w',ZIP_DEFLATED) as target:
        for item in source.infolist():
            data=source.read(item.filename)
            name=item.filename
            if name=='xl/worksheets/sheet1.xml':
                name='xl/custom/sheet.xml'
            elif name=='[Content_Types].xml':
                data=data.replace(b'/xl/worksheets/sheet1.xml',b'/xl/custom/sheet.xml')
            elif name=='xl/_rels/workbook.xml.rels':
                data=data.replace(b'worksheets/sheet1.xml',b'custom/sheet.xml')
            target.writestr(name,data)
    source.close()
    events=list(extract(crafted.getvalue()))
    assert events[0][2]['document_no']=='A-2'


def test_worksheet_relationship_cannot_lie_about_content_type():
    original=BytesIO()
    wb=Workbook()
    ws=wb.active
    ws.append(['Invoice','Date','Department','Customer','Type','Amount','Status'])
    wb.save(original)
    source=ZipFile(BytesIO(original.getvalue()))
    crafted=BytesIO()
    with ZipFile(crafted,'w',ZIP_DEFLATED) as target:
        for item in source.infolist():
            data=source.read(item.filename)
            if item.filename=='[Content_Types].xml':
                data=data.replace(
                    b'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml',
                    b'application/xml',
                )
            target.writestr(item,data)
    source.close()
    try:
        list(extract(crafted.getvalue()))
        assert False, 'worksheet content-type mismatch should be rejected'
    except ValueError as exc:
        assert 'worksheet content type' in str(exc)


def test_malformed_lazy_worksheet_xml_is_a_structured_422(client,auth):
    original=BytesIO()
    wb=Workbook()
    ws=wb.active
    ws.append(['Invoice','Date','Department','Customer','Type','Amount','Status'])
    wb.save(original)
    source=ZipFile(BytesIO(original.getvalue()))
    crafted=BytesIO()
    with ZipFile(crafted,'w',ZIP_DEFLATED) as target:
        for item in source.infolist():
            data=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                data=data.replace(b'</worksheet>',b'')
            target.writestr(item,data)
    source.close()
    payload=crafted.getvalue()
    try:
        list(extract(payload))
        assert False, 'malformed worksheet XML should be rejected'
    except ValueError as exc:
        assert 'Unreadable worksheet XML' in str(exc)
    response=client.post(
        '/api/imports', headers=auth,
        files={'file':('malformed.xlsx',payload,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
    )
    assert response.status_code==422
    assert response.json()['detail']=='Unreadable worksheet XML'


def test_early_gate_does_not_read_unauthorized_body():
    called={'receive':False,'downstream':False}

    async def downstream(scope,receive,send):
        called['downstream']=True

    async def receive():
        called['receive']=True
        return {'type':'http.request','body':b'attack','more_body':False}

    messages=[]
    async def send(message):
        messages.append(message)

    scope={'type':'http','method':'POST','path':'/api/imports','headers':[]}
    asyncio.run(EarlyAPIRequestGate(downstream)(scope,receive,send))
    assert not called['receive'] and not called['downstream']
    assert messages[0]['status']==401
